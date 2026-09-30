"""
Bilingual proofreading script as Excel (.xlsx), and the reverse: read the corrected
script back so the translation can be re-inserted into the same PDF.

Sheet "台本": one row per speech bubble. Only the 译文 column is meant to be edited.
The bubble bbox is kept in hidden columns so rows can be matched to the re-extracted
bubbles even if block ids shift. Sheet "meta" (hidden) records the source PDF hash
and the render parameters of the original run.
"""
import os
import time

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_SHEET = "台本"
META_SHEET = "meta"
HEADERS = ["页码", "气泡", "原文 (OCR)", "译文 (可修改)", "x0", "y0", "x1", "y1"]
COL_TRANSLATION = 4
MATCH_MIN_IOU = 0.5
MAX_XLSX_BYTES = 20 * 1024 * 1024


class CorrectionsError(ValueError):
    """The uploaded Excel is not a proofreading script produced by this tool."""


def _text_cell(ws, row, col, value):
    # Always store as text: a translation starting with "=" must not become a formula
    cell = ws.cell(row=row, column=col)
    cell.value = "" if value is None else str(value)
    cell.data_type = TYPE_STRING
    return cell


class LocalDocumentSkill:
    def __init__(self, output_dir: str):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_bilingual_xlsx(self, task_id: str, pages_data: list, original_filename: str,
                                meta: dict | None = None) -> str:
        """
        pages_data: [{"page_num": 1, "blocks": [{"id": 3, "raw": "...", "translated": "...",
                      "bbox": [x0, y0, x1, y1]}, ...]}, ...]
        meta: extra key/value pairs for the hidden meta sheet (pdf_md5, render params, ...).
        """
        stem = os.path.splitext(original_filename)[0]
        path = os.path.join(self.output_dir, f"script_{task_id[:8]}_{stem}.xlsx")

        wb = Workbook()
        ws = wb.active
        ws.title = SCRIPT_SHEET
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E6F7FF")
        edit_fill = PatternFill("solid", fgColor="FFF7E0")
        wrap = Alignment(wrap_text=True, vertical="top")

        row = 2
        for page in pages_data:
            for blk in page.get("blocks", []):
                if not str(blk.get("raw", "")).strip():
                    continue
                ws.cell(row=row, column=1, value=int(page["page_num"]))
                ws.cell(row=row, column=2, value=int(blk["id"]))
                _text_cell(ws, row, 3, blk.get("raw", "")).alignment = wrap
                t = _text_cell(ws, row, COL_TRANSLATION, blk.get("translated", ""))
                t.alignment = wrap
                t.fill = edit_fill
                for i, v in enumerate(blk.get("bbox") or [0, 0, 0, 0]):
                    ws.cell(row=row, column=5 + i, value=round(float(v), 2))
                row += 1

        ws.freeze_panes = "A2"
        for col, width in zip("ABCD", (7, 7, 45, 45)):
            ws.column_dimensions[col].width = width
        for col in "EFGH":
            ws.column_dimensions[col].hidden = True

        ms = wb.create_sheet(META_SHEET)
        ms.sheet_state = "hidden"
        info = {"source_file": original_filename, "generated": time.strftime("%Y-%m-%d %H:%M:%S")}
        info.update(meta or {})
        for k, v in info.items():
            ms.append([k])
            _text_cell(ms, ms.max_row, 2, v)

        wb.save(path)
        return path


def read_corrections(xlsx_path) -> tuple[dict, dict]:
    """
    `xlsx_path`: a path or a binary file-like object (e.g. io.BytesIO of the upload).
    Returns (meta, corrections) where corrections = {page_num: [{"id", "bbox", "text"}, ...]}.
    Raises CorrectionsError for anything that is not a script from generate_bilingual_xlsx.
    """
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as e:
        raise CorrectionsError(f"无法读取 Excel 文件: {e}")
    try:
        if SCRIPT_SHEET not in wb.sheetnames or META_SHEET not in wb.sheetnames:
            raise CorrectionsError("这不是本系统导出的校对台本（缺少「台本」或 meta 工作表）")
        meta = {}
        for r in wb[META_SHEET].iter_rows(values_only=True):
            if r and r[0]:
                meta[str(r[0])] = "" if len(r) < 2 or r[1] is None else str(r[1])

        rows = wb[SCRIPT_SHEET].iter_rows(values_only=True)
        header = next(rows, None)
        if not header or list(header[:4]) != HEADERS[:4]:
            raise CorrectionsError("「台本」工作表的表头被修改了，请使用原始导出的文件")
        corrections = {}
        for r in rows:
            if not r or r[0] is None or r[1] is None:
                continue
            try:
                page, bid = int(r[0]), int(r[1])
                bbox = [float(v) for v in r[4:8]] if len(r) >= 8 and None not in r[4:8] else None
            except (TypeError, ValueError):
                continue
            text = "" if r[3] is None else str(r[3])
            corrections.setdefault(page, []).append({"id": bid, "bbox": bbox, "text": text})
        return meta, corrections
    finally:
        wb.close()


def _iou(a, b) -> float:
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    union = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / union if union > 0 else 0.0


def match_corrections(blocks: list, rows: list) -> dict:
    """
    Map re-extracted blocks of one page to corrected rows: best bbox IoU (>= MATCH_MIN_IOU),
    same block id breaks ties. Each row is used once. Returns {block_id: text}.
    Blocks without a match get nothing, so the renderer keeps their original text.
    """
    pairs = []
    for bi, blk in enumerate(blocks):
        for ri, row in enumerate(rows):
            if row["bbox"] is not None:
                score = _iou(blk["bbox"], row["bbox"])
                if score < MATCH_MIN_IOU:
                    continue
            elif row["id"] == blk["id"]:
                score = MATCH_MIN_IOU  # no bbox in the sheet: fall back to the id
            else:
                continue
            pairs.append((score, row["id"] == blk["id"], bi, ri))
    pairs.sort(reverse=True)
    used_b, used_r, out = set(), set(), {}
    for _, _, bi, ri in pairs:
        if bi in used_b or ri in used_r:
            continue
        used_b.add(bi)
        used_r.add(ri)
        out[blocks[bi]["id"]] = rows[ri]["text"]
    return out
